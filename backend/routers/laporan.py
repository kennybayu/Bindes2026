"""Ekspor rekap pembayaran sewa lapak ke Excel (openpyxl)."""

import io

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from lib.db import db
from lib.dates import today_iso
from lib.serialize import to_aware
from models.payment import Payment
from routers.auth import require_session

router = APIRouter(prefix="/laporan", tags=["laporan"], dependencies=[Depends(require_session)])

STATUS_LABEL = {"menunggu": "Menunggu", "lunas": "Lunas", "ditolak": "Ditolak"}


def _fmt_dt(value) -> str:
    if not value:
        return "-"
    return value.strftime("%d/%m/%Y %H:%M")


@router.get("/pembayaran.xlsx")
async def export_pembayaran_xlsx():
    docs = await db.payments.find().sort("created_at", -1).to_list(1000)
    payments = [Payment(**to_aware(d, "created_at", "confirmed_at")) for d in docs]

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Pembayaran"

    ws.merge_cells("A1:I1")
    ws["A1"] = "REKAP PEMBAYARAN SEWA LAPAK — PASAR ADAT DESA ADAT JIMBARAN"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A2:I2")
    ws["A2"] = f"Diekspor: {today_iso()} (WITA) • {len(payments)} transaksi"
    ws["A2"].font = Font(size=10, color="64748B")

    headers = [
        "No", "Nomor ID", "Nama Penyewa", "Periode", "Jumlah (Rp)",
        "Metode", "Status", "Tanggal Catat", "Tgl Konfirmasi",
    ]
    header_fill = PatternFill("solid", fgColor="0F172A")
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    total_lunas = 0
    total_menunggu = 0
    row_num = 5
    for idx, p in enumerate(payments, start=1):
        ws.cell(row=row_num, column=1, value=idx)
        ws.cell(row=row_num, column=2, value=p.nomor_id)
        ws.cell(row=row_num, column=3, value=p.nama_lengkap)
        ws.cell(row=row_num, column=4, value=p.periode)
        ws.cell(row=row_num, column=5, value=p.jumlah)
        ws.cell(row=row_num, column=6, value=p.metode.capitalize())
        ws.cell(row=row_num, column=7, value=STATUS_LABEL.get(p.status, p.status))
        ws.cell(row=row_num, column=8, value=_fmt_dt(p.created_at))
        ws.cell(row=row_num, column=9, value=_fmt_dt(p.confirmed_at))
        if p.status == "lunas":
            total_lunas += p.jumlah
        elif p.status == "menunggu":
            total_menunggu += p.jumlah
        row_num += 1

    ws.cell(row=row_num + 1, column=4, value="Total Lunas").font = Font(bold=True)
    ws.cell(row=row_num + 1, column=5, value=total_lunas).font = Font(bold=True)
    ws.cell(row=row_num + 2, column=4, value="Total Menunggu Konfirmasi").font = Font(bold=True)
    ws.cell(row=row_num + 2, column=5, value=total_menunggu).font = Font(bold=True)

    for col, width in enumerate([5, 14, 28, 24, 14, 10, 12, 18, 18], start=1):
        ws.column_dimensions[chr(64 + col)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"rekap-pembayaran-{today_iso()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
